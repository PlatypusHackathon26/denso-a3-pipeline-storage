"""Unit tests for the new ingestion pipeline architecture."""

import io
import tempfile
from pathlib import Path
import docx
import pandas as pd
import pdfplumber
import pytest
from PIL import Image

from configs.settings import settings
from core.hasher import Hasher
from core.router import Router
from core.storage import Storage, storage
from engines.docling_engine import DoclingEngine
from engines.xberg_engine import XbergEngine
from schemas.document import ProcessedDocument
from utils.lang_detector import detect_language, SUPPORTED_LANGS


def test_hasher():
    hasher = Hasher(algorithm="sha256")
    with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tmp:
        tmp.write("hello denso rag")
        tmp_path = Path(tmp.name)

    try:
        h1 = hasher.compute_hash(tmp_path)
        doc_id = hasher.generate_doc_id(tmp_path, h1)
        assert len(h1) == 64
        assert doc_id.startswith("DENSO_")
        assert not hasher.is_processed(h1)

        hasher.register_hash(h1)
        assert hasher.is_processed(h1)
    finally:
        tmp_path.unlink(missing_ok=True)


def test_router():
    router = Router()
    assert isinstance(router.get_engine(Path("report.pdf")), DoclingEngine)
    assert isinstance(router.get_engine(Path("manual.docx")), XbergEngine)
    assert isinstance(router.get_engine(Path("data.xlsx")), XbergEngine)
    assert isinstance(router.get_engine(Path("script.py")), XbergEngine)
    assert router.get_engine(Path("video.mp4")) is None


def test_xberg_text_parser(tmp_path):
    engine = XbergEngine()
    test_file = tmp_path / "sample.txt"
    test_file.write_text("Hướng dẫn xử lý lỗi hệ thống DENSO.", encoding="utf-8")

    doc = engine.parse(test_file, access_level=2)
    assert isinstance(doc, ProcessedDocument)
    assert doc.metadata.source_file == "sample.txt"
    assert doc.metadata.source_type == "txt"
    assert doc.metadata.access_level == 2
    assert "Hướng dẫn xử lý lỗi" in doc.cleaned_text
    assert doc.metadata.language in ("vi", "unknown")


def test_xberg_docx_parser(tmp_path):
    engine = XbergEngine()
    docx_file = tmp_path / "sop_manual.docx"
    doc_obj = docx.Document()
    doc_obj.add_paragraph("Kiểm tra van áp suất P-102")
    tbl = doc_obj.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Thông số"
    tbl.cell(0, 1).text = "Giá trị"
    tbl.cell(1, 0).text = "Áp suất"
    tbl.cell(1, 1).text = "0.5 MPa"
    doc_obj.save(docx_file)

    doc = engine.parse(docx_file, access_level=1)
    assert isinstance(doc, ProcessedDocument)
    assert doc.metadata.source_file == "sop_manual.docx"
    assert doc.metadata.doc_type == "maintenance_manual"
    assert len(doc.tables) == 1
    assert doc.tables[0].table_id.endswith("_t1")
    assert doc.tables[0].total_rows == 2
    assert doc.tables[0].total_cols == 2
    assert "Áp suất" in doc.tables[0].markdown


def test_xberg_excel_parser(tmp_path):
    engine = XbergEngine()

    excel_file = tmp_path / "inventory_data.xlsx"

    df = pd.DataFrame(
        {
            "Linh kiện": ["Cảm biến A", "Bơm dầu B", "Van C"],
            "Mã": ["CB-01", "BD-02", None],
            "Thông số": ["0.5 MPa", None, "45.5 °C"],
        }
    )

    df.to_excel(excel_file, index=False)

    doc = engine.parse(excel_file, access_level=3)

    # 1. Kiểu dữ liệu output
    assert isinstance(doc, ProcessedDocument)

    # 2. Metadata
    assert doc.metadata.source_file == "inventory_data.xlsx"
    assert doc.metadata.source_type == "xlsx"
    assert doc.metadata.access_level == 3

    # 3. Có đúng 1 bảng
    assert len(doc.tables) == 1

    table = doc.tables[0]

    # 4. Thông tin bảng
    assert table.table_id.endswith("_t1")
    assert table.total_rows == 3
    assert table.total_cols == 3

    # 5. Nội dung không bị mất
    assert "Cảm biến A" in table.markdown
    assert "CB-01" in table.markdown
    assert "0.5 MPa" in table.markdown
    assert "45.5 °C" in table.markdown

    # 6. Cell rỗng không được biến thành "nan"
    assert "nan" not in table.markdown.lower()
    assert "nan" not in doc.cleaned_text.lower()

    # 7. Markdown phải là dạng compact
    assert "| --- | --- | --- |" in table.markdown

def test_xberg_excel_multiple_sheets(tmp_path):
    engine = XbergEngine()

    excel_file = tmp_path / "multi_sheet.xlsx"

    error_df = pd.DataFrame(
        {
            "Mã lỗi": ["ERR-305", "ERR-401"],
            "Nguyên nhân": ["Kẹt van", "Quá nhiệt"],
        }
    )

    qc_df = pd.DataFrame(
        {
            "Thông số": ["Áp suất dầu", "Nhiệt độ"],
            "Giá trị": ["0.5 MPa", "45 °C"],
        }
    )

    parameter_df = pd.DataFrame(
        {
            "Tên": ["Voltage", "Rating"],
            "Giá trị": [12, 2.0],
        }
    )

    with pd.ExcelWriter(excel_file) as writer:
        error_df.to_excel(writer, sheet_name="Error Codes", index=False)
        qc_df.to_excel(writer, sheet_name="QC Checklist", index=False)
        parameter_df.to_excel(writer, sheet_name="Parameters", index=False)

    doc = engine.parse(excel_file, access_level=1)

    # Excel có 3 sheet -> phải có 3 table
    assert len(doc.tables) == 3

    # Kiểm tra từng sheet
    assert doc.tables[0].caption == "Bảng: Error Codes"
    assert doc.tables[1].caption == "Bảng: QC Checklist"
    assert doc.tables[2].caption == "Bảng: Parameters"

    # Kiểm tra dữ liệu thực tế
    assert "ERR-305" in doc.tables[0].markdown
    assert "Áp suất dầu" in doc.tables[1].markdown
    assert "Voltage" in doc.tables[2].markdown

def test_xberg_csv_parser(tmp_path):
    engine = XbergEngine()

    csv_file = tmp_path / "error_codes.csv"

    df = pd.DataFrame(
        {
            "Mã lỗi": ["ERR-305", "ERR-401"],
            "Nguyên nhân": ["Kẹt van", "Quá nhiệt"],
            "Cách xử lý": ["Xả van 2 vòng", "Kiểm tra bơm"],
        }
    )

    df.to_csv(csv_file, index=False, encoding="utf-8-sig")

    doc = engine.parse(csv_file, access_level=1)

    assert isinstance(doc, ProcessedDocument)
    assert doc.metadata.source_type == "csv"
    assert len(doc.tables) == 1

    table = doc.tables[0]

    assert table.total_rows == 2
    assert table.total_cols == 3

    assert "ERR-305" in table.markdown
    assert "Kẹt van" in table.markdown
    assert "Xả van 2 vòng" in table.markdown

    assert "nan" not in table.markdown.lower()


def test_storage_save_document(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "PROCESSED_DIR", tmp_path / "processed")
    settings.PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    st = Storage()

    from schemas.document import DocMetadata
    doc = ProcessedDocument(
        doc_id="DENSO_TEST01",
        metadata=DocMetadata(
            source_file="test.txt",
            source_type="txt",
            language="vi",
            access_level=1,
        ),
        cleaned_text="Nội dung test lưu trữ",
    )
    saved_path = st.save_document(doc)
    assert saved_path.exists()

    loaded = st.load_document("test.json")
    assert loaded is not None
    assert loaded.doc_id == "DENSO_TEST01"
    assert loaded.metadata.source_file == "test.txt"
    assert loaded.cleaned_text == "Nội dung test lưu trữ"

def test_docling_table_filter():
    """Bảng rác từ đường kẻ layout phải bị loại, bảng dữ liệu thật phải được giữ."""
    # Không có dữ liệu
    assert DoclingEngine._is_valid_table(None) is False
    assert DoclingEngine._is_valid_table([]) is False

    # Lưới layout rỗng (chỉ có đường kẻ) -> rác
    assert DoclingEngine._is_valid_table([["", "", ""], ["", None, ""]]) is False

    # Khung cảnh báo 1 cột (PRECAUTIONS!/DISCLAIMER) -> không phải bảng dữ liệu
    assert DoclingEngine._is_valid_table([["PRECAUTIONS!"], ["Đọc kỹ trước khi thao tác thiết bị..."]]) is False

    # Khung viền 1 ô
    assert DoclingEngine._is_valid_table([["Hình 1: Sơ đồ hệ thống"]]) is False

    # Bảng có cột nhưng nội dung quá ít (< 10 ký tự) -> rác
    assert DoclingEngine._is_valid_table([["a", "b"], ["c", ""]]) is False

    # Bảng dữ liệu thật
    assert DoclingEngine._is_valid_table(
        [["Thông số", "Giá trị", "Đơn vị"], ["Áp suất", "0.5", "MPa"]]
    ) is True


def test_docling_markdown_table_formatting():
    md = DoclingEngine._format_markdown_table(
        [["Thông số", "Giá trị"], ["Áp suất", "0.5 MPa"], ["Nhiệt độ"]]
    )
    lines = md.splitlines()
    assert lines[0] == "| Thông số | Giá trị |"
    assert lines[1] == "| --- | --- |"
    assert lines[2] == "| Áp suất | 0.5 MPa |"
    # Hàng thiếu ô được pad cho đủ số cột của header
    assert lines[3] == "| Nhiệt độ |  |"

    assert DoclingEngine._format_markdown_table([]) == ""
class _FakeStream:
    """Mô phỏng pdfminer PDFStream: chỉ có get_data(), không có get_raw_data()."""

    def __init__(self, data: bytes) -> None:
        self._data = data

    def get_data(self) -> bytes:
        return self._data


class _FakePage:
    def __init__(self, text, tables=None, images=None, words=None) -> None:
        self._text = text
        self._tables = tables or []
        self.images = images or []
        self._words = words or []

    def extract_text(self) -> str:
        return self._text

    def extract_tables(self):
        return self._tables

    def extract_words(self):
        return self._words


class _FakePdf:
    def __init__(self, pages) -> None:
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *exc_info):
        return False


def _png_bytes(size, mode="L") -> bytes:
    buffer = io.BytesIO()
    Image.new(mode, size, color=128 if mode == "L" else "white").save(buffer, format="PNG")
    return buffer.getvalue()


def test_docling_fallback_extracts_images_and_filters_tables(tmp_path, monkeypatch):
    """Fallback pdfplumber phải lưu ảnh thật và loại bỏ bảng rác/icon nhỏ."""
    page_one = _FakePage(
        text="Hướng dẫn thay thế Suction Control Valve Kit.",
        tables=[
            [["", "", ""], ["", None, ""]],  # lưới layout rỗng -> loại
            [["PRECAUTIONS!"], ["Đọc kỹ hướng dẫn trước khi thao tác."]],  # khung 1 cột -> loại
            [["Thông số", "Giá trị"], ["Áp suất", "0.5 MPa"]],  # bảng thật -> giữ
        ],
        images=[
            {"stream": _FakeStream(_png_bytes((400, 300))), "bottom": 150.0},  # ảnh thật -> giữ
            {"stream": _FakeStream(_png_bytes((20, 20))), "bottom": 300.0},  # icon nhỏ -> loại
            {"stream": _FakeStream(b"khong-phai-anh"), "bottom": 450.0},  # dữ liệu lỗi -> loại
        ],
        # Dòng dưới ảnh không phải caption -> dùng caption mặc định
        words=[{"text": "Installation", "top": 160.0, "x0": 20.0}],
    )
    page_two = _FakePage(
        text="Hình 7: Suction Control Valve Kit.",
        images=[
            {"stream": _FakeStream(_png_bytes((500, 400), mode="RGB")), "bottom": 200.0}
        ],
        # Caption thật nằm ngay dưới ảnh -> ưu tiên dùng caption gốc
        words=[
            {"text": "Hình", "top": 210.0, "x0": 20.0},
            {"text": "7:", "top": 210.0, "x0": 60.0},
            {"text": "Suction", "top": 210.0, "x0": 80.0},
            {"text": "Control", "top": 210.0, "x0": 130.0},
            {"text": "Valve", "top": 210.0, "x0": 180.0},
            {"text": "Kit.", "top": 210.0, "x0": 220.0},
        ],
    )
    monkeypatch.setattr(
        pdfplumber, "open", lambda *args, **kwargs: _FakePdf([page_one, page_two])
    )
    monkeypatch.setattr(storage, "media_dir", tmp_path)
    monkeypatch.setattr(settings, "EXTRACT_IMAGES", True)

    pdf_file = tmp_path / "sop_manual_pdfplumber.pdf"
    pdf_file.write_bytes(b"%PDF-1.4 fake")

    doc = DoclingEngine().parse(pdf_file, access_level=2, doc_id="DENSO_TESTPDF")

    assert doc.metadata.doc_type == "maintenance_manual"
    assert doc.metadata.engine_used.endswith("(pdfplumber fallback)")
    assert doc.metadata.is_corrupted is False
    assert doc.metadata.total_pages == 2
    assert "Hướng dẫn thay thế" in doc.cleaned_text

    # Chỉ 2 ảnh hợp lệ được lưu (icon 20x20 và dữ liệu lỗi bị bỏ)
    assert len(doc.images) == 2

    first_image = doc.images[0]
    assert first_image.image_id == "DENSO_TESTPDF_img1"
    assert first_image.page == 1
    # Không có caption gốc -> sinh caption mặc định kèm context xung quanh
    assert first_image.native_caption == "Hình 1: Trang 1"
    assert first_image.vlm_caption is None
    assert "Hướng dẫn thay thế" in first_image.context_text
    assert (tmp_path / "DENSO_TESTPDF" / first_image.image_name).exists()

    second_image = doc.images[1]
    assert second_image.page == 2
    assert second_image.native_caption == "Hình 7: Suction Control Valve Kit."
    assert (tmp_path / "DENSO_TESTPDF" / second_image.image_name).exists()

    # Chỉ bảng dữ liệu thật được giữ
    assert len(doc.tables) == 1
    assert doc.tables[0].table_id == "DENSO_TESTPDF_t1"
    assert doc.tables[0].total_rows == 2
    assert doc.tables[0].total_cols == 2
    assert "Áp suất" in doc.tables[0].markdown

def test_language_detection_english():
    text = """
    Applications
    Section
    Page
    Brand
    Model
    Engine
    Fuel
    Power
    Application Date
    Denso Part Number
    Voltage
    Rating
    """

    assert detect_language(text) == "en"

def test_language_detection_vietnamese():
    text = """
    Bảng thông số kỹ thuật
    Áp suất dầu
    Nhiệt độ động cơ
    Mã lỗi thiết bị
    """

    assert detect_language(text) == "vi"

def test_language_detection_supported_language_only():
    text = """
    Applications Brand Model Engine Fuel Power
    Denso Part Number Voltage Rating
    """

    result = detect_language(text)

    assert result in SUPPORTED_LANGS